"""يبني نموذج دراسة الجدوى (قالب الهيئة العامة للعقار - البيع على الخارطة)
من بيانات ملف "نموذج ارض العليا" مع ربط كل الأرقام بالمعادلات.

الخلايا الزرقاء = مدخلات رقمية، الخلايا الصفراء = بيانات مطلوبة من العميل.
"""
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName

OUT = "output/Feasibility_Study_Olaya.xlsx"

TEAL = "0F7C74"
LIGHT = "E3F1EF"
GREY = "F2F2F2"
YELLOW = "FFF2CC"
FONT = "Arial"

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
NUM = '#,##0;(#,##0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'

PERIODS = ["فترة التجهيزات", "السنة الأولى", "السنة الثانية", "السنة الثالثة"]
PCOLS = ["D", "E", "F", "G"]  # نفس أعمدة الفترات في المدخلات والحسابات

wb = Workbook()
wb.remove(wb.active)


def sheet(title):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    return ws


def style(c, bold=False, fill=None, color="000000", fmt=None, size=10, align=None, border=True):
    c.font = Font(name=FONT, bold=bold, color=color, size=size)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if fmt:
        c.number_format = fmt
    if border:
        c.border = BORDER
    c.alignment = Alignment(horizontal=align or "right", vertical="center", wrap_text=True)


def title(ws, text, sub=None):
    ws["B1"] = text
    style(ws["B1"], bold=True, color=TEAL, size=16, border=False)
    if sub:
        ws["B2"] = sub
        style(ws["B2"], color="7F7F7F", size=9, border=False)


def section(ws, row, text, span=6):
    ws.cell(row, 2, text)
    for col in range(2, 2 + span):
        style(ws.cell(row, col), bold=True, fill=TEAL, color="FFFFFF", size=11)


def header(ws, row, labels, start=2):
    for i, t in enumerate(labels):
        c = ws.cell(row, start + i, t)
        style(c, bold=True, fill=LIGHT, align="center")


def name(n, ws, cell):
    col = "".join(ch for ch in cell if ch.isalpha())
    row = "".join(ch for ch in cell if ch.isdigit())
    wb.defined_names[n] = DefinedName(n, attr_text=f"'{ws.title}'!${col}${row}")


def ref(ws, cell):
    return f"'{ws.title}'!{cell}"


# ---------------------------------------------------------------- المدخلات
inp = sheet("المدخلات")
title(inp, "المدخلات والافتراضات", "المصدر: نموذج ارض العليا - بتاريخ 11-08-2026 | الأزرق = مدخل رقمي، الأصفر = مطلوب من العميل")
for col, w in zip("BCDEFG", [48, 22, 18, 18, 18, 18]):
    inp.column_dimensions[col].width = w

TEXT_INPUTS = [
    ("ProjectName", "اسم المشروع", "مشروع أرض العليا - شقق سكنية"),
    ("StudyTitle", "عنوان الدراسة", "دراسة جدوى مشروع شقق سكنية للبيع على الخارطة"),
    ("City", "المدينة", None),
    ("District", "الحي", "العليا"),
    ("Streets", "الشوارع الرئيسية", None),
    ("DeedNo", "رقم الصك", None),
    ("UrbanScope", "نطاق الموقع", None),
    ("LandUse", "الاستخدامات", "سكني"),
    ("Developer", "المطور العقاري", None),
    ("Consultant", "المستشار الهندسي الرئيسي", None),
    ("Auditor", "المحاسب القانوني", None),
    ("StartDate", "التاريخ المتوقع لبداية المشروع", None),
    ("EndDate", "التاريخ المتوقع لانتهاء المشروع", None),
    ("HandoverDate", "التاريخ المتوقع لتسليم الوحدات", None),
]

NUM_INPUTS = [
    ("وصف المشروع والمساحات", None, None, None),
    ("LandArea", "مساحة أرض المشروع (م2)", 3000, NUM),
    ("BuildRatio", "نسبة البناء على الأرض", 0.5, PCT),
    ("FAR", "معامل البناء", 2.37166666666667, "0.0000"),
    ("NetRatio", "نسبة المساحة الصافية (القابلة للبيع)", 0.85, PCT),
    ("UnitArea", "متوسط مساحة الوحدة (م2)", 250, NUM),
    ("ParkPerUnit", "متطلبات المواقف لكل وحدة", 5, NUM2),
    ("ParkArea", "متوسط مساحة الموقف بالقبو (م2)", 25, NUM),
    ("تكاليف الأرض", None, None, None),
    ("LandPrice", "سعر شراء متر الأرض (ريال)", 3200, NUM),
    ("LandFeeRate", "السعي والضريبة (% من قيمة الأرض)", 0.075, PCT),
    ("تكاليف البناء", None, None, None),
    ("ResCostM2", "تكلفة بناء المتر - السكني (ريال)", 1500, NUM),
    ("BasementCostM2", "تكلفة بناء المتر - القبو والمواقف (ريال)", 600, NUM),
    ("OtherCostRate", "التكاليف الأخرى: تصاميم، إشراف، رسوم حكومية، تسويق (% من التكاليف المباشرة)", 0.02, PCT),
    ("ContRate", "الاحتياطي (% من التكاليف المباشرة)", 0, PCT),
    ("DevFeeRate", "أتعاب المطور (% من التكاليف المباشرة)", 0, PCT),
    ("التمويل البنكي", None, None, None),
    ("LoanAmount", "مبلغ التمويل البنكي (ريال)", 8700000, NUM),
    ("MurabahaRate", "نسبة المرابحة السنوية", 0.07, PCT),
    ("LoanFeeRate", "رسوم ترتيب التمويل", 0.01, PCT),
    ("أخرى", None, None, None),
    ("Months", "مدة المشروع (شهر)", 22, NUM),
    ("ZakatRate", "نسبة الزكاة / الضريبة على صافي الربح", 0, PCT),
]

r = 4
section(inp, r, "أولاً | بيانات المشروع (نصية)", 2)
r += 1
for key, label, val in TEXT_INPUTS:
    inp.cell(r, 2, label)
    style(inp.cell(r, 2))
    c = inp.cell(r, 3, val)
    style(c, fill=YELLOW if val is None else None, color="1F4E79")
    name(key, inp, f"C{r}")
    r += 1

r += 1
section(inp, r, "ثانياً | الافتراضات الرقمية", 2)
r += 1
for key, label, val, fmt in NUM_INPUTS:
    if label is None:
        inp.cell(r, 2, key)
        style(inp.cell(r, 2), bold=True, fill=GREY)
        style(inp.cell(r, 3), fill=GREY)
    else:
        inp.cell(r, 2, label)
        style(inp.cell(r, 2))
        style(inp.cell(r, 3, val), color="0000FF", fmt=fmt)
        name(key, inp, f"C{r}")
    r += 1

r += 1
section(inp, r, "ثالثاً | الجداول الزمنية (نسب التوزيع على الفترات)", 6)
r += 1
header(inp, r, ["البند", "الإجمالي"] + PERIODS)
r += 1
SCHEDULES = [
    ("LandPct", "نسبة سداد تكلفة الأرض", [1, 0, 0, 0], PCT),
    ("ConstrPct", "نسبة صرف تكاليف البناء والتطوير", [0, 0.5, 0.5, 0], PCT),
    ("SalesPct", "نسبة المساحة المباعة", [0, 0.5, 0.5, 0], PCT),
    ("PriceM2", "متوسط سعر بيع المتر (ريال)", [0, 4400, 5400, 5400], NUM),
    ("DrawPct", "نسبة سحب التمويل البنكي", [0, 0.5, 0.5, 0], PCT),
    ("RepayPct", "نسبة سداد التمويل البنكي", [0, 0, 1, 0], PCT),
]
SCH = {}
for key, label, vals, fmt in SCHEDULES:
    inp.cell(r, 2, label)
    style(inp.cell(r, 2))
    if fmt == PCT:
        style(inp.cell(r, 3, f"=SUM(D{r}:G{r})"), bold=True, fmt=PCT)
    else:
        style(inp.cell(r, 3))
    for col, v in zip(PCOLS, vals):
        style(inp[f"{col}{r}"], color="0000FF", fmt=fmt)
        inp[f"{col}{r}"] = v
    SCH[key] = r
    r += 1
inp.cell(r + 1, 2, "ملاحظة: التحصيل النقدي من المبيعات يتم 100% في سنة البيع (كما في الملف الأصلي).")
style(inp.cell(r + 1, 2), color="7F7F7F", size=9, border=False)


def sch(key, col):
    return f"'المدخلات'!{col}{SCH[key]}"


# ---------------------------------------------------------------- الحسابات
calc = sheet("الحسابات")
title(calc, "الحسابات التفصيلية", "لا تحتوي على مدخلات - كل الخلايا معادلات مرتبطة بورقة المدخلات")
for col, w in zip("BCDEFG", [46, 18, 18, 18, 18, 18]):
    calc.column_dimensions[col].width = w

SCALARS = [
    ("المساحات", None, None),
    ("Footprint", "المساحة المبنى عليها (م2)", "=LandArea*BuildRatio"),
    ("BUA", "المسطحات البنائية (م2)", "=LandArea*FAR"),
    ("Floors", "عدد الأدوار", "=IFERROR(BUA/Footprint,0)"),
    ("NetArea", "المساحات الصافية القابلة للبيع (م2)", "=BUA*NetRatio"),
    ("Units", "عدد الوحدات", "=IFERROR(NetArea/UnitArea,0)"),
    ("Parking", "عدد المواقف المطلوبة", "=Units*ParkPerUnit"),
    ("BasementArea", "مساحة القبو - مواقف (م2)", "=Parking*ParkArea"),
    ("BasementFloors", "عدد أدوار القبو", "=IFERROR(BasementArea/LandArea,0)"),
    ("التكاليف", None, None),
    ("LandPurchase", "قيمة شراء الأرض", "=LandArea*LandPrice"),
    ("LandFees", "السعي والضريبة", "=LandPurchase*LandFeeRate"),
    ("LandCost", "إجمالي تكلفة الأرض", "=LandPurchase+LandFees"),
    ("ResCost", "تكاليف المباني والإنشاءات (السكني)", "=BUA*ResCostM2"),
    ("BasementCost", "تكاليف البنية التحتية والمرافق (القبو والمواقف)", "=BasementArea*BasementCostM2"),
    ("DirectCost", "التكاليف المباشرة", "=ResCost+BasementCost"),
    ("OtherCost", "التكاليف الأخرى (تصاميم، إشراف، رسوم حكومية، تسويق)", "=DirectCost*OtherCostRate"),
    ("ContCost", "الاحتياطي", "=DirectCost*ContRate"),
    ("DevFee", "أتعاب المطور", "=DirectCost*DevFeeRate"),
    ("IndirectCost", "التكاليف غير المباشرة", "=OtherCost+ContCost+DevFee"),
    ("DevCost", "إجمالي تكاليف التطوير (بدون الأرض)", "=DirectCost+IndirectCost"),
    ("ProjectCost", "إجمالي تكلفة المشروع (قبل التمويل)", "=LandCost+DevCost"),
    ("LoanFee", "رسوم ترتيب التمويل", "=LoanAmount*LoanFeeRate"),
]
r = 4
section(calc, r, "أولاً | الحسابات الأساسية", 2)
r += 1
for key, label, f in SCALARS:
    if label is None:
        calc.cell(r, 2, key)
        style(calc.cell(r, 2), bold=True, fill=GREY)
        style(calc.cell(r, 3), fill=GREY)
    else:
        calc.cell(r, 2, label)
        style(calc.cell(r, 2))
        style(calc.cell(r, 3, f), fmt=NUM2 if key in ("Floors", "Units", "Parking", "BasementFloors") else NUM)
        name(key, calc, f"C{r}")
    r += 1

r += 1
section(calc, r, "ثانياً | الحسابات حسب الفترات", 6)
r += 1
header(calc, r, ["البند", "الإجمالي"] + PERIODS)
r += 1

# (key, label, formula template, total: "sum" | None | template, fmt, bold)
# {c}=عمود الفترة، {p}=العمود السابق، {R[key]}=رقم صف بند آخر
PROWS = [
    ("_", "الإيرادات والتحصيل", None, None, None, True),
    ("SoldArea", "المساحات المباعة (م2)", "={s}*NetArea", "sum", NUM, False),
    ("UnitsSold", "عدد الوحدات المباعة", "={s}*Units", "sum", NUM2, False),
    ("Price", "متوسط سعر المتر (ريال)", "={pr}", "=IFERROR(C{R[Revenue]}/C{R[SoldArea]},0)", NUM, False),
    ("UnitPrice", "متوسط سعر الوحدة (ريال)", "={c}{R[Price]}*UnitArea", "=C{R[Price]}*UnitArea", NUM, False),
    ("Revenue", "إيرادات المبيعات", "={c}{R[SoldArea]}*{c}{R[Price]}", "sum", NUM, True),
    ("_", "التدفقات الخارجة", None, None, None, True),
    ("LandOut", "تكاليف الأرض", "={lp}*LandCost", "sum", NUM, False),
    ("ResOut", "تكاليف المباني والإنشاءات", "={cp}*ResCost", "sum", NUM, False),
    ("BasOut", "تكاليف البنية التحتية والمرافق (القبو)", "={cp}*BasementCost", "sum", NUM, False),
    ("OtherOut", "التكاليف الأخرى (تصاميم، إشراف، رسوم، تسويق)", "={cp}*OtherCost", "sum", NUM, False),
    ("ContOut", "الاحتياطي", "={cp}*ContCost", "sum", NUM, False),
    ("DevFeeOut", "أتعاب المطور", "={cp}*DevFee", "sum", NUM, False),
    ("_", "التمويل البنكي", None, None, None, True),
    ("LoanOpen", "رصيد القرض أول المدة", "={prev_loan}", None, NUM, False),
    ("Draw", "سحب", "={dp}*LoanAmount", "sum", NUM, False),
    ("Repay", "سداد", "={rp}*LoanAmount", "sum", NUM, False),
    ("LoanClose", "رصيد القرض آخر المدة", "={c}{R[LoanOpen]}+{c}{R[Draw]}-{c}{R[Repay]}", None, NUM, False),
    ("Murabaha", "المرابحة", "=({c}{R[LoanOpen]}+{c}{R[Draw]})*MurabahaRate", "sum", NUM, False),
    ("LoanFeeOut", "رسوم ترتيب التمويل", "={first}*LoanFee", "sum", NUM, False),
    ("_", "قائمة الدخل (مطابقة التكاليف مع الإيرادات)", None, None, None, True),
    ("RevShare", "نسبة الإيراد المعترف به", "=IFERROR({c}{R[Revenue]}/$C{R[Revenue]},0)", "sum", PCT, False),
    ("IS_Land", "تكلفة الأرض المباعة", "={c}{R[RevShare]}*LandCost", "sum", NUM, False),
    ("IS_Res", "تكاليف المباني والإنشاءات", "={c}{R[RevShare]}*ResCost", "sum", NUM, False),
    ("IS_Bas", "تكاليف البنية التحتية والمرافق", "={c}{R[RevShare]}*BasementCost", "sum", NUM, False),
    ("IS_Gross", "مجمل الربح", "={c}{R[Revenue]}-{c}{R[IS_Land]}-{c}{R[IS_Res]}-{c}{R[IS_Bas]}", "sum", NUM, True),
    ("IS_Other", "التكاليف الأخرى (تصاميم، إشراف، رسوم، تسويق)", "={c}{R[RevShare]}*OtherCost", "sum", NUM, False),
    ("IS_Cont", "الاحتياطي وأتعاب المطور", "={c}{R[RevShare]}*(ContCost+DevFee)", "sum", NUM, False),
    ("IS_EBF", "صافي الربح قبل رسوم التمويل والزكاة", "={c}{R[IS_Gross]}-{c}{R[IS_Other]}-{c}{R[IS_Cont]}", "sum", NUM, True),
    ("IS_Fin", "رسوم وتكاليف التمويل", "={c}{R[Murabaha]}+{c}{R[LoanFeeOut]}", "sum", NUM, False),
    ("IS_PBZ", "صافي الربح قبل الزكاة", "={c}{R[IS_EBF]}-{c}{R[IS_Fin]}", "sum", NUM, True),
    ("Zakat", "الزكاة / الضريبة", "=MAX(0,{c}{R[IS_PBZ]})*ZakatRate", "sum", NUM, False),
    ("NetProfit", "صافي الربح", "={c}{R[IS_PBZ]}-{c}{R[Zakat]}", "sum", NUM, True),
    ("_", "التدفقات النقدية", None, None, None, True),
    ("CashOut", "إجمالي التدفقات الخارجة", "={c}{R[LandOut]}+{c}{R[ResOut]}+{c}{R[BasOut]}+{c}{R[OtherOut]}+{c}{R[ContOut]}+{c}{R[DevFeeOut]}+{c}{R[Murabaha]}+{c}{R[LoanFeeOut]}+{c}{R[Zakat]}", "sum", NUM, True),
    ("NetOps", "صافي التدفق قبل التمويل", "={c}{R[Revenue]}-{c}{R[CashOut]}", "sum", NUM, False),
    ("NetBE", "صافي التدفق بعد التمويل البنكي (قبل الملاك)", "={c}{R[NetOps]}+{c}{R[Draw]}-{c}{R[Repay]}", "sum", NUM, True),
    ("CashOpen", "رصيد النقدية أول المدة", "={prev_cash}", None, NUM, False),
    ("Equity", "ضخ الملاك / المساهمين (النقد المطلوب)", "=MAX(0,-({c}{R[CashOpen]}+{c}{R[NetBE]}))", "sum", NUM, False),
    ("CashClose", "رصيد النقدية آخر المدة", "={c}{R[CashOpen]}+{c}{R[NetBE]}+{c}{R[Equity]}", None, NUM, True),
]

# الترقيم أولاً ثم الكتابة (حتى تعمل المراجع الأمامية)
R = {}
rr = r
for row in PROWS:
    if row[0] != "_":
        R[row[0]] = rr
    rr += 1

for key, label, tmpl, total, fmt, bold in PROWS:
    calc.cell(r, 2, label)
    if key == "_":
        for col in range(2, 8):
            style(calc.cell(r, col), bold=True, fill=GREY)
        r += 1
        continue
    style(calc.cell(r, 2), bold=bold)
    for i, c in enumerate(PCOLS):
        p = PCOLS[i - 1] if i else None
        f = tmpl.format(
            c=c, R=R,
            s=sch("SalesPct", c), pr=sch("PriceM2", c), lp=sch("LandPct", c),
            cp=sch("ConstrPct", c), dp=sch("DrawPct", c), rp=sch("RepayPct", c),
            first=1 if i == 0 else 0,
            prev_loan=f"{p}{R['LoanClose']}" if p else 0,
            prev_cash=f"{p}{R['CashClose']}" if p else 0,
        )
        style(calc.cell(r, i + 4, f), bold=bold, fmt=fmt)
    if total == "sum":
        style(calc.cell(r, 3, f"=SUM(D{r}:G{r})"), bold=True, fmt=fmt, fill=LIGHT)
    elif total:
        style(calc.cell(r, 3, total.format(R=R)), bold=True, fmt=fmt, fill=LIGHT)
    else:
        style(calc.cell(r, 3), fill=LIGHT)
    r += 1

for key in ["Revenue", "LandOut", "ResOut", "BasOut", "OtherOut", "Murabaha", "LoanFeeOut",
            "Equity", "NetProfit", "Zakat", "IS_Gross", "IS_EBF", "IS_Fin", "IS_PBZ", "Draw", "UnitsSold",
            "SoldArea"]:
    name("T_" + key, calc, f"C{R[key]}")

r += 1
section(calc, r, "ثالثاً | مؤشرات الربحية", 2)
r += 1
KPIS = [
    ("TotalInvest", "حجم الاستثمار (تكلفة المشروع + التمويل)", "=ProjectCost+T_Murabaha+T_LoanFeeOut", NUM),
    ("BuyersFund", "تمويل من متحصلات المشترين", "=TotalInvest-T_Equity-LoanAmount", NUM),
    ("ROI_Total", "العائد الكلي على رأس مال الملاك", "=IFERROR(T_NetProfit/T_Equity,0)", PCT),
    ("ROI", "معدل العائد السنوي على الاستثمار (ROI)", "=IFERROR(ROI_Total/(Months/12),0)", PCT),
    ("IRR", "معدل العائد الداخلي (IRR)", f"=IFERROR(IRR(D{R['NetBE']}:G{R['NetBE']}),0)", PCT),
    ("MarginSales", "هامش صافي الربح من المبيعات", "=IFERROR(T_NetProfit/T_Revenue,0)", PCT),
    ("MarginCost", "صافي الربح إلى إجمالي التكاليف", "=IFERROR(T_NetProfit/TotalInvest,0)", PCT),
    ("CostPerUnit", "متوسط تكلفة الوحدة", "=IFERROR(TotalInvest/Units,0)", NUM),
]
for key, label, f, fmt in KPIS:
    calc.cell(r, 2, label)
    style(calc.cell(r, 2))
    style(calc.cell(r, 3, f), bold=True, fmt=fmt)
    name(key, calc, f"C{r}")
    r += 1


def cp(key, col):
    """مرجع لخلية فترة في الحسابات"""
    return f"='الحسابات'!{col}{R[key]}"


# ------------------------------------------------------- أوراق قالب الهيئة
def kv_table(ws, row, rows, fmt=NUM, widths=None):
    for label, f, *rest in rows:
        f_fmt = rest[0] if rest else fmt
        ws.cell(row, 2, label)
        style(ws.cell(row, 2))
        c = ws.cell(row, 3, f)
        if f is None:
            style(c, fill=YELLOW)
        else:
            style(c, fmt=f_fmt, bold=True)
        row += 1
    return row


def placeholder(ws, row, label, lines=3):
    ws.cell(row, 2, label)
    style(ws.cell(row, 2), bold=True, fill=LIGHT)
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=7)
    style(ws.cell(row, 3), fill=YELLOW)
    ws.row_dimensions[row].height = 18 * lines
    return row + 1


# الملخص التنفيذي --------------------------------------------------------
ex = sheet("الملخص التنفيذي")
title(ex, "الملخص التنفيذي")
for col, w in zip("BCDEFG", [46, 22, 16, 16, 16, 16]):
    ex.column_dimensions[col].width = w
r = 3
kv_table(ex, r, [("عنوان الدراسة", "=StudyTitle", "@"), ("موقع المشروع", '=District&" - "&City', "@")])
r = 6
section(ex, r, "ملخص الدراسة المالية", 2)
r += 1
header(ex, r, ["العناصر", "القيمة (ريال)"])
r += 1
SUMMARY = [
    ("صافي المبيعات - الوحدات السكنية", "=T_Revenue"),
    ("إجمالي الإيرادات", "=T_Revenue"),
    ("إجمالي تكاليف الأرض الخام", "=LandCost"),
    ("تكاليف الإنشاءات", "=ResCost"),
    ("تكاليف تطوير البنية التحتية (القبو والمواقف)", "=BasementCost"),
    ("مصاريف الاستشاري الهندسي", 0),
    ("مصاريف المحاسب القانوني", 0),
    ("تكاليف الإشراف والمتابعة والعمولة", 0),
    ("تكاليف التسويق", 0),
    ("تكاليف الرواتب والأجور", 0),
    ("مصاريف إدارية وعمومية", 0),
    ("تكاليف التمويل", "=T_Murabaha+T_LoanFeeOut"),
    ("تكاليف ورسوم أخرى (تصاميم، إشراف، رسوم حكومية، تسويق - مجمعة)", "=IndirectCost"),
    ("إجمالي التكاليف الإنشائية والفنية", None),
    ("الزكاة / الضريبة", "=T_Zakat"),
    ("صافي الربح", "=T_NetProfit"),
]
first = r
for label, f in SUMMARY:
    ex.cell(r, 2, label)
    bold = label.startswith("إجمالي") or label == "صافي الربح"
    if label == "إجمالي التكاليف الإنشائية والفنية":
        f = f"=SUM(C{first + 3}:C{r - 1})"
    style(ex.cell(r, 2), bold=bold, fill=LIGHT if bold else None)
    style(ex.cell(r, 3, f), fmt=NUM, bold=bold, fill=LIGHT if bold else None)
    r += 1
ex.cell(r, 2, "* الملف الأصلي يجمع مصاريف الاستشاري والتصميم والإشراف والتسويق والرسوم في بند واحد (2%) - يلزم التفصيل.")
style(ex.cell(r, 2), color="C00000", size=9, border=False)

r += 2
section(ex, r, "مؤشرات التقييم / الجدوى", 2)
r += 1
r = kv_table(ex, r, [
    ("تكاليف الأرض (ريال)", "=LandCost"),
    ("تكاليف التطوير (ريال)", "=DevCost"),
    ("معدل العائد على الاستثمار (ROI) - سنوي", "=ROI", PCT),
    ("معدل العائد الداخلي - العائد على التطوير (IRR)", "=IRR", PCT),
])
r += 1
section(ex, r, "ملخص الدراسة التمهيدية والتسويقية", 6)
r = placeholder(ex, r + 1, "الملخص", 4)
section(ex, r, "ملخص الدراسة الفنية", 6)
r += 1
ex.cell(r, 2, "=\"مشروع شقق سكنية على أرض مساحتها \"&TEXT(LandArea,\"#,##0\")&\" م2، بمسطحات بناء \"&TEXT(BUA,\"#,##0\")&\" م2 على \"&TEXT(Floors,\"0.0\")&\" أدوار تقريباً، ومساحة صافية للبيع \"&TEXT(NetArea,\"#,##0\")&\" م2 (\"&TEXT(Units,\"0\")&\" وحدة تقريباً بمتوسط \"&TEXT(UnitArea,\"0\")&\" م2)، مع قبو مواقف بمساحة \"&TEXT(BasementArea,\"#,##0\")&\" م2.\"")
ex.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
style(ex.cell(r, 2))
ex.row_dimensions[r].height = 48
r += 2
section(ex, r, "الفترة الزمنية للتطوير والفحص والتسليم", 2)
r = kv_table(ex, r + 1, [
    ("بداية المشروع", "=IF(StartDate=\"\",\"\",StartDate)", "yyyy/mm/dd"),
    ("مدة تنفيذ المشروع (شهر)", "=Months"),
    ("نهاية المشروع", "=IF(EndDate=\"\",\"\",EndDate)", "yyyy/mm/dd"),
])

chart_row = 7
ex["E7"], ex["F7"] = "البند", "مليون ريال"
chart_items = [("إجمالي الإيرادات", "=T_Revenue/1e6"), ("تكاليف الأرض", "=LandCost/1e6"),
               ("التكاليف الإنشائية والفنية", "=(DevCost+T_Murabaha+T_LoanFeeOut)/1e6"),
               ("صافي الربح", "=T_NetProfit/1e6")]
for i, (l, f) in enumerate(chart_items):
    ex.cell(8 + i, 5, l)
    ex.cell(8 + i, 6, f).number_format = "0.0"
    style(ex.cell(8 + i, 5))
    style(ex.cell(8 + i, 6), fmt="0.0")
header(ex, 7, ["البند", "مليون ريال"], start=5)
ch = BarChart()
ch.title = "ملخص مالي (مليون ريال)"
ch.legend = None
ch.add_data(Reference(ex, min_col=6, min_row=7, max_row=11), titles_from_data=True)
ch.set_categories(Reference(ex, min_col=5, min_row=8, max_row=11))
ch.series[0].graphicalProperties.solidFill = TEAL
ch.height, ch.width = 7, 13
ex.add_chart(ch, "E13")

# الدراسة التمهيدية والتسويقية -------------------------------------------
mk = sheet("الدراسة التسويقية")
title(mk, "الدراسة التمهيدية والتسويقية", "الخلايا الصفراء بيانات نوعية مطلوبة من العميل")
for col, w in zip("BCDEFGHIJ", [34, 22, 16, 16, 16, 16, 14, 14, 14]):
    mk.column_dimensions[col].width = w
r = 4
section(mk, r, "نبذة عامة للمشروع", 6)
r += 1
mk.cell(r, 2, "وصف فكرة المشروع")
style(mk.cell(r, 2), bold=True, fill=LIGHT)
mk.cell(r, 3, "=\"تطوير \"&TEXT(Units,\"0\")&\" شقة سكنية تقريباً بمتوسط مساحة \"&UnitArea&\" م2 في حي \"&District&\" وبيعها على الخارطة خلال \"&Months&\" شهراً.\"")
mk.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
style(mk.cell(r, 3))
mk.row_dimensions[r].height = 36
r += 1
for label, f in [("المطور العقاري", "=IF(Developer=\"\",\"\",Developer)"),
                 ("نوع وعدد ومساحات الوحدات السكنية", "=\"شقق سكنية - \"&TEXT(Units,\"0.0\")&\" وحدة - متوسط \"&UnitArea&\" م2\""),
                 ("المستشار الهندسي الرئيسي", "=IF(Consultant=\"\",\"\",Consultant)"),
                 ("المحاسب القانوني (إذا انطبق)", "=IF(Auditor=\"\",\"\",Auditor)")]:
    mk.cell(r, 2, label)
    style(mk.cell(r, 2), bold=True, fill=LIGHT)
    mk.cell(r, 3, f)
    mk.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
    style(mk.cell(r, 3))
    r += 1
r += 1
section(mk, r, "حجم القطاع العقاري - مساهمة القطاع في الناتج المحلي (آخر 3 سنوات)", 8)
r += 1
header(mk, r, ["الفئة / السنة", "2023 قيمة الناتج", "% من الناتج", "معدل النمو", "2024 قيمة الناتج", "% من الناتج", "معدل النمو"])
for i, lab in enumerate(["الأنشطة العقارية", "الناتج المحلي الإجمالي"]):
    mk.cell(r + 1 + i, 2, lab)
    style(mk.cell(r + 1 + i, 2))
    for col in range(3, 9):
        style(mk.cell(r + 1 + i, col), fill=YELLOW)
r += 3
r = placeholder(mk, r, "نبذة عن الوضع الحالي والتطورات المستقبلية للقطاع", 3)
r += 1
section(mk, r, "العملاء المستهدفين", 6)
r = placeholder(mk, r + 1, "شرائح العملاء")
r = placeholder(mk, r, "قنوات الوصول للعملاء")
r = placeholder(mk, r, "العلاقات مع العملاء")
r += 1
section(mk, r, "المشاريع المنافسة", 8)
r += 1
header(mk, r, ["اسم المشروع", "المساحة (م2)", "نوع الوحدات", "عدد الوحدات", "متوسط سعر المتر", "الموقع", "المطور"])
for i in range(5):
    for col in range(2, 9):
        style(mk.cell(r + 1 + i, col), fill=YELLOW)
r += 7
section(mk, r, "التحليل الرباعي (SWOT)", 6)
for lab in ["نقاط القوة", "نقاط الضعف", "الفرص", "التهديدات"]:
    r = placeholder(mk, r + 1, lab) - 1
r += 2
section(mk, r, "المزيج التسويقي", 6)
r += 1
mk.cell(r, 2, "الخدمات / المنتجات")
style(mk.cell(r, 2), bold=True, fill=LIGHT)
mk.cell(r, 3, "=\"شقق سكنية بمتوسط مساحة \"&UnitArea&\" م2 مع مواقف في القبو\"")
mk.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
style(mk.cell(r, 3))
r += 1
mk.cell(r, 2, "التسعير")
style(mk.cell(r, 2), bold=True, fill=LIGHT)
mk.cell(r, 3, f"=\"متوسط سعر المتر \"&TEXT('المدخلات'!E{SCH['PriceM2']},\"#,##0\")&\" ريال في المرحلة الأولى و\"&TEXT('المدخلات'!F{SCH['PriceM2']},\"#,##0\")&\" ريال في المرحلة الثانية (متوسط \"&TEXT('الحسابات'!C{R['Price']},\"#,##0\")&\" ريال)\"")
mk.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
style(mk.cell(r, 3))
r = placeholder(mk, r + 1, "الترويج")
r = placeholder(mk, r, "التوزيع")

# الدراسة الفنية ---------------------------------------------------------
te = sheet("الدراسة الفنية")
title(te, "الدراسة الفنية")
for col, w in zip("BCDEFG", [42, 20, 18, 20, 18, 18]):
    te.column_dimensions[col].width = w
r = 4
section(te, r, "تحليل الموقع", 2)
r = kv_table(te, r + 1, [
    ("موقع المشروع (حي - مدينة)", '=District&" - "&City', "@"),
    ("الشوارع الرئيسية", "=IF(Streets=\"\",\"\",Streets)", "@"),
    ("نطاق الموقع", "=IF(UrbanScope=\"\",\"\",UrbanScope)", "@"),
    ("المساحة (وفقاً للصك) م2", "=LandArea"),
    ("رقم الصك", "=IF(DeedNo=\"\",\"\",DeedNo)", "@"),
    ("الاستخدامات", "=LandUse", "@"),
    ("المساحة السكنية م2", "=BUA"),
    ("المساحة التجارية م2", 0),
    ("المساحة السكنية التجارية م2", 0),
])
r = placeholder(te, r, "نبذة عن موقع المشروع وإحداثياته (خريطة)", 3)
r += 1
section(te, r, "التوزيع المساحي", 6)
r += 1
header(te, r, ["البيان", "المساحة (م2)", "النسبة من الأرض"])
r += 1
for label, f in [("مساحة المخطط بالكامل (الأرض)", "=LandArea"),
                 ("المساحة المبنى عليها (الأرضي)", "=Footprint"),
                 ("المسطحات البنائية فوق الأرض", "=BUA"),
                 ("المساحة الصافية القابلة للبيع", "=NetArea"),
                 ("مساحة القبو (مواقف)", "=BasementArea")]:
    te.cell(r, 2, label)
    style(te.cell(r, 2))
    style(te.cell(r, 3, f), fmt=NUM)
    style(te.cell(r, 4, f"=IFERROR(C{r}/LandArea,0)"), fmt=PCT)
    r += 1
r = kv_table(te, r, [("عدد الأدوار", "=Floors", "0.00"), ("عدد أدوار القبو", "=BasementFloors", "0.00"),
                     ("عدد المواقف المطلوبة", "=Parking", "0")])
r += 1
section(te, r, "نبذة عن الوحدات العقارية", 6)
r += 1
header(te, r, ["نوع الوحدات", "متوسط المساحة م2", "عدد الوحدات", "متوسط سعر التعاقد (ريال)", "تاريخ التسليم المتوقع"])
r += 1
for i, c in enumerate(PCOLS[1:3]):
    te.cell(r, 2, f"شقق سكنية - {PERIODS[i + 1]}")
    style(te.cell(r, 2))
    style(te.cell(r, 3, "=UnitArea"), fmt=NUM)
    style(te.cell(r, 4, cp("UnitsSold", c)), fmt=NUM2)
    style(te.cell(r, 5, cp("UnitPrice", c)), fmt=NUM)
    style(te.cell(r, 6, "=IF(HandoverDate=\"\",\"\",HandoverDate)"), fmt="yyyy/mm/dd")
    r += 1
r += 1
section(te, r, "المدة الزمنية", 2)
r = kv_table(te, r + 1, [
    ("التاريخ المتوقع لبداية المشروع", "=IF(StartDate=\"\",\"\",StartDate)", "yyyy/mm/dd"),
    ("التاريخ المتوقع لانتهاء المشروع", "=IF(EndDate=\"\",\"\",EndDate)", "yyyy/mm/dd"),
    ("التاريخ المتوقع لتسليم الوحدات", "=IF(HandoverDate=\"\",\"\",HandoverDate)", "yyyy/mm/dd"),
    ("مدة المشروع (شهر)", "=Months"),
])
r += 1
section(te, r, "عناصر التكاليف الاستثمارية", 4)
r += 1
header(te, r, ["تكاليف الأرض", "المساحة م2", "سعر المتر", "القيمة الإجمالية"])
r += 1
for label, a, b, f in [("قيمة شراء الأرض", "=LandArea", "=LandPrice", "=LandPurchase"),
                       ("السعي والضريبة", None, "=LandFeeRate", "=LandFees"),
                       ("إجمالي تكاليف الأرض", "=LandArea", "=IFERROR(LandCost/LandArea,0)", "=LandCost")]:
    te.cell(r, 2, label)
    style(te.cell(r, 2), bold=label.startswith("إجمالي"))
    style(te.cell(r, 3, a), fmt=NUM)
    style(te.cell(r, 4, b), fmt=PCT if b == "=LandFeeRate" else NUM)
    style(te.cell(r, 5, f), fmt=NUM, bold=label.startswith("إجمالي"))
    r += 1
r += 1
header(te, r, ["التكاليف الإنشائية", "المساحة م2", "تكلفة الإنشاء بالمتر", "القيمة الإجمالية"])
r += 1
for label, a, b, f in [("تكاليف المباني والإنشاءات", "=BUA", "=ResCostM2", "=ResCost"),
                       ("تكاليف البنية التحتية والمرافق (القبو والمواقف)", "=BasementArea", "=BasementCostM2", "=BasementCost"),
                       ("القيمة الإجمالية لتكاليف الإنشاءات", "=BUA+BasementArea", "=IFERROR(DirectCost/(BUA+BasementArea),0)", "=DirectCost")]:
    te.cell(r, 2, label)
    style(te.cell(r, 2), bold=label.startswith("القيمة"))
    style(te.cell(r, 3, a), fmt=NUM)
    style(te.cell(r, 4, b), fmt=NUM)
    style(te.cell(r, 5, f), fmt=NUM, bold=label.startswith("القيمة"))
    r += 1
r += 1
header(te, r, ["التكاليف والمصاريف الإدارية والفنية", "القيمة الإجمالية"])
r += 1
first = r
for label, f in [("مصاريف الاستشاري الهندسي", 0), ("مصاريف التصميم", 0), ("مصاريف المحاسب القانوني", 0),
                 ("تكاليف الإشراف والمتابعة والعمولة", 0), ("تكاليف التسويق", 0), ("تكاليف إدارية وعمومية", 0),
                 ("تكاليف الرواتب والأجور", 0), ("تكاليف التمويل (مرابحة + رسوم ترتيب)", "=T_Murabaha+T_LoanFeeOut"),
                 ("تكاليف أخرى (تصاميم، إشراف، رسوم حكومية، تسويق - مجمعة)", "=OtherCost"),
                 ("الاحتياطي وأتعاب المطور", "=ContCost+DevFee")]:
    te.cell(r, 2, label)
    style(te.cell(r, 2))
    style(te.cell(r, 3, f), fmt=NUM)
    r += 1
te.cell(r, 2, "القيمة الإجمالية للتكاليف الإدارية والفنية")
style(te.cell(r, 2), bold=True, fill=LIGHT)
style(te.cell(r, 3, f"=SUM(C{first}:C{r - 1})"), bold=True, fmt=NUM, fill=LIGHT)
r += 2
te.cell(r, 2, "إجمالي قيمة التكاليف الاستثمارية المتوقعة")
style(te.cell(r, 2), bold=True, fill=TEAL, color="FFFFFF")
style(te.cell(r, 3, "=TotalInvest"), bold=True, fmt=NUM, fill=TEAL, color="FFFFFF")

# الدراسة المالية --------------------------------------------------------
fi = sheet("الدراسة المالية")
title(fi, "الدراسة المالية")
for col, w in zip("BCDEFGH", [52, 18, 18, 18, 18, 18, 18]):
    fi.column_dimensions[col].width = w
r = 4
section(fi, r, "الإيرادات المتوقعة", 6)
r += 1
header(fi, r, ["نوع الوحدة", "المساحة المباعة م2", "عدد الوحدات", "متوسط سعر المتر", "متوسط سعر الوحدة", "القيمة الإجمالية"])
r += 1
first = r
for i, c in enumerate(PCOLS[1:]):
    fi.cell(r, 2, f"الوحدات السكنية - {PERIODS[i + 1]}")
    style(fi.cell(r, 2))
    for j, (k, fmt) in enumerate([("SoldArea", NUM), ("UnitsSold", NUM2), ("Price", NUM), ("UnitPrice", NUM), ("Revenue", NUM)]):
        style(fi.cell(r, 3 + j, cp(k, c)), fmt=fmt)
    r += 1
fi.cell(r, 2, "إجمالي قيمة الإيرادات المتوقعة")
style(fi.cell(r, 2), bold=True, fill=LIGHT)
for j, (k, fmt) in enumerate([("SoldArea", NUM), ("UnitsSold", NUM2), ("Price", NUM), ("UnitPrice", NUM), ("Revenue", NUM)]):
    style(fi.cell(r, 3 + j, cp(k, "C")), fmt=fmt, bold=True, fill=LIGHT)
r += 2

section(fi, r, "الهيكل التمويلي", 6)
r += 1
header(fi, r, ["العناصر", "القيمة الإجمالية", "النسبة المئوية", "المرحلة", "عدد الدفعات"])
r += 1
first = r
for label, f, stage, n in [("حقوق الملكية (الملاك / المساهمين)", "=T_Equity", "فترة التجهيزات (شراء الأرض)", "=COUNTIF('الحسابات'!D%d:G%d,\">0\")" % (R["Equity"], R["Equity"])),
                           ("قروض (التمويل البنكي)", "=LoanAmount", "السنة الأولى والثانية", "=COUNTIF('المدخلات'!D%d:G%d,\">0\")" % (SCH["DrawPct"], SCH["DrawPct"])),
                           ("مشترين (متحصلات البيع على الخارطة)", "=BuyersFund", "خلال فترة البناء", None)]:
    fi.cell(r, 2, label)
    style(fi.cell(r, 2))
    style(fi.cell(r, 3, f), fmt=NUM)
    style(fi.cell(r, 4, f"=IFERROR(C{r}/C{first + 3},0)"), fmt=PCT)
    style(fi.cell(r, 5, stage))
    style(fi.cell(r, 6, n), fill=YELLOW if n is None else None, align="center")
    r += 1
fi.cell(r, 2, "إجمالي التمويل")
style(fi.cell(r, 2), bold=True, fill=LIGHT)
style(fi.cell(r, 3, f"=SUM(C{first}:C{r - 1})"), bold=True, fmt=NUM, fill=LIGHT)
style(fi.cell(r, 4, f"=SUM(D{first}:D{r - 1})"), bold=True, fmt=PCT, fill=LIGHT)
style(fi.cell(r, 5), fill=LIGHT)
style(fi.cell(r, 6), fill=LIGHT)
r += 2


def period_table(ws, r, title_text, rows):
    section(ws, r, title_text, 6)
    r += 1
    header(ws, r, ["البند"] + PERIODS + ["الإجمالي"])
    r += 1
    for label, key, kind in rows:
        ws.cell(r, 2, label)
        bold = kind == "bold"
        neg = kind == "neg"
        style(ws.cell(r, 2), bold=bold, fill=LIGHT if bold else None)
        for i, c in enumerate(PCOLS + ["C"]):
            if key is None:
                f = None
            elif isinstance(key, str) and key.startswith("="):
                f = key.replace("{c}", c)
            else:
                f = ("=-" if neg else "=") + f"'الحسابات'!{c}{R[key]}"
            if key in ("CashOpen", "CashClose") and c == "C":
                f = None
            fmt = PCT if key == "_pct" else NUM
            style(ws.cell(r, 3 + i, f), fmt=fmt, bold=bold, fill=LIGHT if bold else None)
        r += 1
    return r


r = period_table(fi, r, "قائمة الدخل (الأرباح والخسائر)", [
    ("إجمالي الإيرادات", "Revenue", "bold"),
    ("تكاليف الأرض", "IS_Land", "neg"),
    ("تكاليف المباني والإنشاءات", "IS_Res", "neg"),
    ("تكاليف البنية التحتية والمرافق", "IS_Bas", "neg"),
    ("مجمل الربح", "IS_Gross", "bold"),
    ("تكاليف أخرى (تصاميم، إشراف، رسوم حكومية، تسويق)", "IS_Other", "neg"),
    ("الاحتياطي وأتعاب المطور", "IS_Cont", "neg"),
    ("صافي الربح قبل رسوم التمويل والزكاة", "IS_EBF", "bold"),
    ("رسوم التمويل", "IS_Fin", "neg"),
    ("صافي الربح قبل الزكاة أو الضريبة", "IS_PBZ", "bold"),
    ("الزكاة أو الضريبة", "Zakat", "neg"),
    ("صافي الربح أو الخسارة", "NetProfit", "bold"),
])
# نسبة صافي الربح
fi.cell(r, 2, "نسبة صافي الربح أو الخسارة")
style(fi.cell(r, 2), bold=True, fill=LIGHT)
for i, c in enumerate(PCOLS + ["C"]):
    style(fi.cell(r, 3 + i, f"=IFERROR('الحسابات'!{c}{R['NetProfit']}/'الحسابات'!{c}{R['Revenue']},0)"), fmt=PCT, bold=True, fill=LIGHT)
r += 1
fi.cell(r, 2, "* الإيراد يُعترف به في سنة البيع، وتُحمّل تكاليف الأرض والبناء بنسبة المبيعات (مبدأ المقابلة)؛ تكاليف التمويل حسب حدوثها.")
style(fi.cell(r, 2), color="7F7F7F", size=9, border=False)
r += 2

r = period_table(fi, r, "قائمة التدفقات النقدية", [
    ("التدفقات النقدية الداخلة", None, "bold"),
    ("المبيعات - الوحدات السكنية", "Revenue", ""),
    ("إجمالي التدفقات النقدية الداخلة", "Revenue", "bold"),
    ("التدفقات النقدية الخارجة", None, "bold"),
    ("تكاليف الأرض", "LandOut", "neg"),
    ("تكاليف المباني والإنشاءات", "ResOut", "neg"),
    ("تكاليف البنية التحتية والمرافق", "BasOut", "neg"),
    ("تكاليف أخرى (تصاميم، إشراف، رسوم حكومية، تسويق)", "OtherOut", "neg"),
    ("الاحتياطي وأتعاب المطور", "='الحسابات'!{c}%d*-1-'الحسابات'!{c}%d" % (R["ContOut"], R["DevFeeOut"]), ""),
    ("رسوم التمويل (مرابحة + ترتيب)", "='الحسابات'!{c}%d*-1-'الحسابات'!{c}%d" % (R["Murabaha"], R["LoanFeeOut"]), ""),
    ("الزكاة أو الضريبة", "Zakat", "neg"),
    ("إجمالي التدفقات النقدية الخارجة", "CashOut", "neg"),
    ("صافي التدفق النقدي التشغيلي", "NetOps", "bold"),
    ("التدفقات التمويلية", None, "bold"),
    ("سحب التمويل البنكي", "Draw", ""),
    ("سداد التمويل البنكي", "Repay", "neg"),
    ("ضخ الملاك / المساهمين", "Equity", ""),
    ("صافي رصيد النقدية في بداية الفترة", "CashOpen", ""),
    ("صافي رصيد النقدية في نهاية الفترة", "CashClose", "bold"),
])

# النتائج والتوصيات -------------------------------------------------------
rs = sheet("النتائج والتوصيات")
title(rs, "نتائج وتوصيات الدراسة")
for col, w in zip("BCDEFG", [40, 40, 16, 16, 16, 16]):
    rs.column_dimensions[col].width = w
r = 4
section(rs, r, "النتائج التقديرية للمشروع", 2)
r += 1
header(rs, r, ["البند", "القيمة"])
r = kv_table(rs, r + 1, [
    ("إجمالي الإيرادات - القيمة الإجمالية للإيرادات المتوقعة (ريال)", "=T_Revenue"),
    ("إجمالي التكاليف - القيمة الإجمالية للتكاليف (ريال)", "=TotalInvest+T_Zakat"),
    ("صافي الأرباح - القيمة الإجمالية لصافي الأرباح (ريال)", "=T_NetProfit"),
    ("رأس مال الملاك المطلوب (ريال)", "=T_Equity"),
    ("العائد الكلي على رأس المال", "=ROI_Total", PCT),
    ("معدل العائد على الاستثمار ROI (سنوي)", "=ROI", PCT),
    ("معدل العائد الداخلي IRR (التطوير)*", "=IRR", PCT),
    ("هامش صافي الربح من المبيعات", "=MarginSales", PCT),
    ("متوسط تكلفة الوحدة (ريال)", "=CostPerUnit"),
    ("متوسط سعر بيع الوحدة (ريال)", f"='الحسابات'!C{R['UnitPrice']}"),
    ("نسبة المرابحة البنكية (للمقارنة)", "=MurabahaRate", PCT),
])
rs.cell(r, 2, "* يقارن معدل العائد الداخلي بسعر الفائدة السائد (مع افتراض ثبات باقي العوامل الأخرى).")
style(rs.cell(r, 2), color="7F7F7F", size=9, border=False)
r += 2
section(rs, r, "التوصيات", 6)
r += 1
rs.cell(r, 2, '=IF(IRR>MurabahaRate,"المشروع مجدٍ مالياً: معدل العائد الداخلي ("&TEXT(IRR,"0.0%")&") أعلى من تكلفة التمويل ("&TEXT(MurabahaRate,"0.0%")&")، وصافي الربح المتوقع "&TEXT(T_NetProfit,"#,##0")&" ريال بعائد سنوي "&TEXT(ROI,"0.0%")&" على رأس مال الملاك.","المشروع غير مجدٍ بالافتراضات الحالية: معدل العائد الداخلي أقل من تكلفة التمويل.")')
rs.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
style(rs.cell(r, 2), bold=True)
rs.row_dimensions[r].height = 48
r += 1
for lab in ["التوصيات السوقية", "التوصيات الفنية", "التوصيات المالية"]:
    rs.cell(r, 2, lab)
    style(rs.cell(r, 2), bold=True, fill=LIGHT)
    rs.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
    style(rs.cell(r, 3), fill=YELLOW)
    rs.row_dimensions[r].height = 40
    r += 1

# ترتيب الأوراق حسب القالب
order = ["الملخص التنفيذي", "الدراسة التسويقية", "الدراسة الفنية", "الدراسة المالية",
         "النتائج والتوصيات", "المدخلات", "الحسابات"]
wb._sheets = [wb[n] for n in order]
wb.active = 0
for ws in wb:
    ws.sheet_properties.tabColor = TEAL if ws.title not in ("المدخلات", "الحسابات") else "808080"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0

from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print("saved", OUT)
