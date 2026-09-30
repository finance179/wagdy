"""نموذج جدوى مبسط لإقناع مستثمر بشراء 12 تاون هاوس في مشروع القادسية تون هاوس
بسعر 1,150,000 ريال للوحدة، مع مقارنة بأسعار المشاريع والعروض المحيطة.

الخلايا الزرقاء = مدخلات يمكن تعديلها، الصفراء = افتراض يحتاج تأكيد.
"""
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties

OUT = "output/Investor_Pitch_Qadisiyah_Townhouses.xlsx"

TEAL = "0F7C74"
LIGHT = "E3F1EF"
GREY = "F2F2F2"
YELLOW = "FFF2CC"
GREEN = "E2EFDA"
FONT = "Arial"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
NUM = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'

wb = Workbook()
wb.remove(wb.active)


def sheet(title):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    return ws


def style(c, bold=False, fill=None, color="000000", fmt=None, size=10, align=None, border=True):
    c.font = Font(name=FONT, bold=bold, color=color, size=size)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if fmt:
        c.number_format = fmt
    if border:
        c.border = BORDER
    c.alignment = Alignment(horizontal=align or "right", vertical="center", wrap_text=True)


def section(ws, row, text, span):
    ws.cell(row, 2, text)
    for col in range(2, 2 + span):
        style(ws.cell(row, col), bold=True, fill=TEAL, color="FFFFFF", size=11)


def header(ws, row, labels, start=2):
    for i, t in enumerate(labels):
        style(ws.cell(row, start + i, t), bold=True, fill=LIGHT, align="center")


def name(n, ws, cell):
    col = "".join(ch for ch in cell if ch.isalpha())
    row = "".join(ch for ch in cell if ch.isdigit())
    wb.defined_names[n] = DefinedName(n, attr_text=f"'{ws.title}'!${col}${row}")


# ------------------------------------------------------------------ المدخلات
inp = sheet("المدخلات")
inp["B1"] = "المدخلات والافتراضات"
style(inp["B1"], bold=True, color=TEAL, size=16, border=False)
inp["B2"] = "الأزرق = مدخل قابل للتعديل | الأصفر = افتراض يحتاج تأكيد"
style(inp["B2"], color="7F7F7F", size=9, border=False)
inp.column_dimensions["B"].width = 58
inp.column_dimensions["C"].width = 20
inp.column_dimensions["D"].width = 60

INPUTS = [
    ("الصفقة", None, None, None, None),
    ("Units", "عدد الوحدات", 12, NUM, None),
    ("UnitPrice", "سعر الوحدة للمستثمر (ريال)", 1150000, NUM, None),
    ("UnitArea", "مساحة الوحدة (م2)", 250, NUM, "افتراض من ملف القادسية (متوسط مساحة الوحدة 250 م2) - يرجى التأكيد"),
    ("Months", "مدة الاحتفاظ حتى التسليم وإعادة البيع (شهر)", 22, NUM, "مدة المشروع في ملف القادسية"),
    ("DownPct", "نسبة الدفع عند التعاقد (الباقي عند التسليم)", 1, PCT, "100% = دفع كامل مقدم (أكثر تحفظاً)"),
    ("تكاليف الشراء والبيع", None, None, None, None),
    ("BuyCostRate", "تكاليف الشراء: ضريبة التصرفات العقارية", 0.05, PCT, "تُصفَّر لو تحملها المطور كحافز للصفقة"),
    ("SellCostRate", "تكاليف إعادة البيع: سعي / تسويق", 0.025, PCT, None),
    ("سيناريوهات إعادة البيع (خصم عن متوسط سعر المتر في السوق)", None, None, None, None),
    ("DiscCons", "المتحفظ", 0.15, PCT, "البيع بأقل 15% من متوسط السوق"),
    ("DiscBase", "الأساسي", 0.10, PCT, "البيع بأقل 10% من متوسط السوق"),
    ("DiscOpt", "المتفائل", 0.0, PCT, "البيع بمتوسط السوق"),
    ("سيناريو التأجير", None, None, None, None),
    ("Rent", "الإيجار السنوي للوحدة (ريال)", 85000, NUM, "عروض القادسية 72-85 ألف لمساحة 173-189 م2؛ وحدتنا أكبر"),
    ("RentCostRate", "مصاريف التشغيل والصيانة والإدارة (% من الإيجار)", 0.10, PCT, None),
    ("للمقارنة", None, None, None, None),
    ("BenchRate", "عائد بديل (ودائع / صكوك) سنوياً", 0.05, PCT, None),
]
r = 4
for key, label, val, fmt, note in INPUTS:
    if label is None:
        inp.cell(r, 2, key)
        for col in (2, 3, 4):
            style(inp.cell(r, col), bold=True, fill=GREY)
    else:
        inp.cell(r, 2, label)
        style(inp.cell(r, 2))
        yellow = key in ("UnitArea", "BuyCostRate", "Rent")
        style(inp.cell(r, 3, val), color="0000FF", fmt=fmt, fill=YELLOW if yellow else None)
        name(key, inp, f"C{r}")
        style(inp.cell(r, 4, note), color="7F7F7F", size=9)
    r += 1

# ------------------------------------------------------------ المشاريع المحيطة
cm = sheet("المشاريع المحيطة")
cm["B1"] = "أسعار التاون هاوس في القادسية والأحياء المجاورة (شرق الرياض)"
style(cm["B1"], bold=True, color=TEAL, size=16, border=False)
cm["B2"] = "عروض معلنة على المنصات العقارية (أسعار طلب وليست صفقات منفذة) - تم البحث في سبتمبر 2026"
style(cm["B2"], color="7F7F7F", size=9, border=False)
for col, w in zip("BCDEFGH", [16, 40, 12, 16, 14, 14, 22]):
    cm.column_dimensions[col].width = w

HARAJ_950 = "https://haraj.com.sa/11114895518/"
HARAJ_1200 = "https://haraj.com.sa/11151387860/"
HARAJ_Q = "https://haraj.com.sa/11140216303/"
HARAJ_YAR = "https://haraj.com.sa/en/11166554307/"
PF_YAR = "https://www.propertyfinder.sa/ar/plp/%D9%84%D9%84%D8%A8%D9%8A%D8%B9/%D8%B4%D9%82%D8%A9-%D9%84%D9%84%D8%A8%D9%8A%D8%B9-%D8%A7%D9%84%D8%B1%D9%8A%D8%A7%D8%B6-%D8%A7%D9%84%D8%B1%D9%8A%D8%A7%D8%B6-%D8%A7%D9%84%D9%8A%D8%B1%D9%85%D9%88%D9%83-255167.html"
AQAR_NAHDA = "https://sa.aqar.fm/%D9%81%D9%84%D9%84-%D9%84%D9%84%D8%A8%D9%8A%D8%B9/%D8%A7%D9%84%D8%B1%D9%8A%D8%A7%D8%B6/%D8%B4%D8%B1%D9%82-%D8%A7%D9%84%D8%B1%D9%8A%D8%A7%D8%B6/%D8%AD%D9%8A-%D8%A7%D9%84%D9%86%D9%87%D8%B6%D8%A9"

COMPS = [
    ("القادسية", "تاون هاوس", 200, 950900, "حراج", HARAJ_950),
    ("القادسية", "فيلا تاون هاوس", 220, 1200000, "حراج", HARAJ_1200),
    ("القادسية", "تاون هاوس 3 أدوار - عمر سنتين", 188, 1300000, "حراج", HARAJ_Q),
    ("اليرموك", "تاون هاوس فاخر 4 غرف", 251.51, 1800000, "حراج", HARAJ_YAR),
    ("اليرموك", "تاون هاوس", 253.5, 1550000, "بروبرتي فايندر", PF_YAR),
    ("النهضة", "تاون هاوس", 264, 1500000, "عقار", AQAR_NAHDA),
    ("النهضة", "تاون هاوس علوي", 180, 1200000, "عقار", AQAR_NAHDA),
]
r = 4
section(cm, r, "أولاً | عروض بيع تاون هاوس (مساحة معلنة)", 7)
r += 1
header(cm, r, ["الحي", "الوصف", "المساحة م2", "السعر (ريال)", "سعر المتر", "سعرنا أقل بـ", "المصدر"])
r += 1
first = r
for dist, desc, area, price, src, url in COMPS:
    for col, v in zip(range(2, 6), [dist, desc, area, price]):
        style(cm.cell(r, col, v), fmt=NUM if col >= 4 else None)
    style(cm.cell(r, 6, f"=E{r}/D{r}"), fmt=NUM)
    style(cm.cell(r, 7, f"=1-OurPriceM2/F{r}"), fmt=PCT)
    c = cm.cell(r, 8, src)
    c.hyperlink = url
    style(c, color="0563C1")
    r += 1
last = r - 1
cm.cell(r, 2, "مشروعنا - القادسية تون هاوس")
style(cm.cell(r, 2), bold=True, fill=GREEN)
style(cm.cell(r, 3, "تاون هاوس جديد - بيع بالجملة 12 وحدة"), bold=True, fill=GREEN)
style(cm.cell(r, 4, "=UnitArea"), bold=True, fill=GREEN, fmt=NUM)
style(cm.cell(r, 5, "=UnitPrice"), bold=True, fill=GREEN, fmt=NUM)
style(cm.cell(r, 6, "=OurPriceM2"), bold=True, fill=GREEN, fmt=NUM)
style(cm.cell(r, 7), fill=GREEN)
style(cm.cell(r, 8), fill=GREEN)
ours_row = r
r += 2
for label, f, key in [("متوسط سعر المتر في السوق", f"=AVERAGE(F{first}:F{last})", "MktAvgM2"),
                      ("وسيط سعر المتر في السوق", f"=MEDIAN(F{first}:F{last})", "MktMedM2"),
                      ("متوسط سعر الوحدة المعروضة", f"=AVERAGE(E{first}:E{last})", "MktAvgUnit")]:
    cm.cell(r, 2, label)
    cm.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    style(cm.cell(r, 2), bold=True, fill=LIGHT)
    style(cm.cell(r, 6, f), bold=True, fmt=NUM, fill=LIGHT)
    name(key, cm, f"F{r}")
    r += 1
r += 1

section(cm, r, "ثانياً | عروض أخرى في الحي (بدون مساحة معلنة / للمقارنة)", 7)
r += 1
header(cm, r, ["الحي", "الوصف", "المساحة م2", "السعر (ريال)", "", "", "المصدر"])
r += 1
for dist, desc, area, price, src, url in [
    ("القادسية", "تاون هاوس علوي", "-", 1200000, "حراج", "https://haraj.com.sa/en/11164555558/"),
    ("القادسية", "تاون هاوس أرضي", "-", 1500000, "حراج", HARAJ_Q),
    ("القادسية", "تاون هاوس علوي", "-", 1400000, "حراج", HARAJ_Q),
    ("اليرموك", "دور تاون هاوس أرضي", "-", 1750000, "حراج", HARAJ_YAR),
    ("القادسية", "شقة - مشروع ركايا (ركز)", 181, 956000, "ركز", "https://rakez.sa/en/project/1880/"),
    ("القادسية", "شقة / بنتهاوس - مشروع نورفا (ركز)", 273, 1125000, "ركز", "https://rakez.sa/en/project/30814/"),
]:
    for col, v in zip(range(2, 6), [dist, desc, area, price]):
        style(cm.cell(r, col, v), fmt=NUM if col >= 4 else None)
    style(cm.cell(r, 6))
    style(cm.cell(r, 7))
    c = cm.cell(r, 8, src)
    c.hyperlink = url
    style(c, color="0563C1")
    r += 1
cm.cell(r, 2, "ملاحظة: شقق مشروع نورفا في نفس الحي تبدأ من 1,125,000 ريال - أي أن سعر التاون هاوس عندنا قريب من سعر الشقة.")
style(cm.cell(r, 2), color="C00000", size=9, border=False)
r += 2

section(cm, r, "ثالثاً | إيجارات التاون هاوس في القادسية", 7)
r += 1
header(cm, r, ["الحي", "الوصف", "المساحة م2", "الإيجار السنوي", "", "", "المصدر"])
r += 1
for dist, desc, area, rent, src, url in [
    ("القادسية", "تاون هاوس دورين - مجمع ركايا (دفعة واحدة)", 173, 75000, "حراج", "https://haraj.com.sa/11156104033/"),
    ("القادسية", "تاون هاوس دورين - مجمع ركايا (دفعتين)", 173, 85000, "حراج", "https://haraj.com.sa/11156104033/"),
    ("القادسية", "تاون هاوس دورين 3 غرف", 189, "-", "عقار (X)", "https://x.com/aqarapp/status/1892597582922195433"),
    ("القادسية", "تاون هاوس (دفعة واحدة)", "-", 72000, "حراج", "https://haraj.com.sa/11156104033/"),
]:
    for col, v in zip(range(2, 6), [dist, desc, area, rent]):
        style(cm.cell(r, col, v), fmt=NUM if col >= 4 else None)
    style(cm.cell(r, 6))
    style(cm.cell(r, 7))
    c = cm.cell(r, 8, src)
    c.hyperlink = url
    style(c, color="0563C1")
    r += 1

# ------------------------------------------------------------- الحسابات (مخفية الوظيفة داخل العرض)
calc = sheet("الحسابات")
calc["B1"] = "الحسابات"
style(calc["B1"], bold=True, color=TEAL, size=16, border=False)
calc.column_dimensions["B"].width = 50
for col in "CDE":
    calc.column_dimensions[col].width = 18

r = 3
section(calc, r, "الصفقة", 2)
r += 1
for key, label, f in [
    ("OurPriceM2", "سعر المتر للمستثمر", "=UnitPrice/UnitArea"),
    ("DealValue", "قيمة الصفقة (12 وحدة)", "=Units*UnitPrice"),
    ("BuyCosts", "تكاليف الشراء", "=DealValue*BuyCostRate"),
    ("TotalCost", "إجمالي تكلفة المستثمر", "=DealValue+BuyCosts"),
    ("Discount", "خصم سعرنا عن متوسط السوق", "=1-OurPriceM2/MktAvgM2"),
    ("MktValue", "القيمة السوقية للوحدات بمتوسط السوق", "=Units*UnitArea*MktAvgM2"),
    ("InstantEquity", "فرق القيمة لصالح المستثمر يوم الشراء", "=MktValue-DealValue"),
    ("CashIn", "النقد المستثمر فعلياً حتى التسليم (دفعة التعاقد + تكاليف الشراء)", "=DealValue*DownPct+BuyCosts"),
]:
    calc.cell(r, 2, label)
    style(calc.cell(r, 2))
    style(calc.cell(r, 3, f), fmt=PCT if key == "Discount" else NUM, bold=True)
    name(key, calc, f"C{r}")
    r += 1

r += 1
section(calc, r, "سيناريوهات إعادة البيع عند التسليم", 4)
r += 1
header(calc, r, ["البند", "متحفظ", "أساسي", "متفائل"])
r += 1
SC = {}
for key, label, tmpl, fmt in [
    ("PM2", "سعر بيع المتر", "=MktAvgM2*(1-{d})", NUM),
    ("PU", "سعر بيع الوحدة", "={c}{PM2}*UnitArea", NUM),
    ("Gross", "إجمالي المبيعات (12 وحدة)", "={c}{PU}*Units", NUM),
    ("SellC", "تكاليف البيع", "={c}{Gross}*SellCostRate", NUM),
    ("Net", "صافي المتحصلات", "={c}{Gross}-{c}{SellC}", NUM),
    ("Profit", "صافي الربح", "={c}{Net}-TotalCost", NUM),
    ("ROI", "العائد الكلي على إجمالي التكلفة", "={c}{Profit}/TotalCost", PCT),
    ("ROIa", "العائد السنوي المركب", "=(1+{c}{ROI})^(12/Months)-1", PCT),
    ("ROIc", "العائد على النقد المستثمر فعلياً (حسب نسبة الدفع المقدم)", "={c}{Profit}/CashIn", PCT),
    ("ROIca", "العائد السنوي على النقد المستثمر", "=(1+{c}{ROIc})^(12/Months)-1", PCT),
    ("Break", "سعر المتر اللازم للتعادل", "=TotalCost/(1-SellCostRate)/Units/UnitArea", NUM),
]:
    SC[key] = r
    calc.cell(r, 2, label)
    style(calc.cell(r, 2), bold=key in ("Profit", "ROIa"))
    for c, d in zip("CDE", ["DiscCons", "DiscBase", "DiscOpt"]):
        f = tmpl.format(c=c, d=d, **{k: v for k, v in SC.items()})
        style(calc.cell(r, "CDE".index(c) + 3, f), fmt=fmt, bold=key in ("Profit", "ROIa"))
    r += 1

r += 1
section(calc, r, "سيناريو التأجير", 2)
r += 1
for key, label, f, fmt in [
    ("RentGross", "إجمالي الإيجار السنوي (12 وحدة)", "=Rent*Units", NUM),
    ("RentNet", "صافي الإيجار بعد المصاريف", "=RentGross*(1-RentCostRate)", NUM),
    ("YieldGross", "العائد الإيجاري الإجمالي", "=RentGross/TotalCost", PCT),
    ("YieldNet", "العائد الإيجاري الصافي", "=RentNet/TotalCost", PCT),
    ("Payback", "فترة استرداد رأس المال من الإيجار (سنة)", "=TotalCost/RentNet", "0.0"),
    ("MktYield", "العائد الإيجاري لو اشترى بسعر السوق", "=RentNet/(MktValue*(1+BuyCostRate))", PCT),
]:
    calc.cell(r, 2, label)
    style(calc.cell(r, 2))
    style(calc.cell(r, 3, f), fmt=fmt, bold=True)
    name(key, calc, f"C{r}")
    r += 1

# ------------------------------------------------------------------ عرض المستثمر
ex = sheet("عرض المستثمر")
for col, w in zip("BCDEFG", [44, 18, 18, 18, 4, 30]):
    ex.column_dimensions[col].width = w
ex["B1"] = "فرصة استثمارية: 12 تاون هاوس - القادسية تون هاوس (شرق الرياض)"
style(ex["B1"], bold=True, color=TEAL, size=16, border=False)
ex["B2"] = "دراسة جدوى مبسطة لشراء 12 وحدة بسعر جملة مقارنةً بأسعار السوق المحيط"
style(ex["B2"], color="7F7F7F", size=10, border=False)

r = 4
section(ex, r, "الصفقة المقترحة", 3)
r += 1
for label, f, fmt in [
    ("عدد الوحدات", "=Units", NUM),
    ("سعر الوحدة (ريال)", "=UnitPrice", NUM),
    ("مساحة الوحدة (م2)", "=UnitArea", NUM),
    ("قيمة الصفقة (ريال)", "=DealValue", NUM),
    ("إجمالي التكلفة شاملة ضريبة التصرفات (ريال)", "=TotalCost", NUM),
]:
    ex.cell(r, 2, label)
    style(ex.cell(r, 2))
    ex.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    style(ex.cell(r, 3, f), fmt=fmt, bold=True)
    style(ex.cell(r, 4))
    r += 1

r += 1
section(ex, r, "لماذا هذه الصفقة؟ - مقارنة بالسوق", 3)
r += 1
for label, f, fmt, hl in [
    ("سعر المتر لنا", "=OurPriceM2", NUM, True),
    ("متوسط سعر المتر للتاون هاوس في القادسية والجوار", "=MktAvgM2", NUM, False),
    ("خصم سعرنا عن متوسط السوق", "=Discount", PCT, True),
    ("القيمة السوقية للـ12 وحدة بمتوسط السوق", "=MktValue", NUM, False),
    ("فرق القيمة لصالح المستثمر من يوم الشراء", "=InstantEquity", NUM, True),
]:
    ex.cell(r, 2, label)
    style(ex.cell(r, 2), bold=hl, fill=GREEN if hl else None)
    ex.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    style(ex.cell(r, 3, f), fmt=fmt, bold=True, fill=GREEN if hl else None)
    style(ex.cell(r, 4), fill=GREEN if hl else None)
    r += 1

r += 1
section(ex, r, "الخيار (1): إعادة البيع عند التسليم", 4)
r += 1
header(ex, r, ["البند", "متحفظ", "أساسي", "متفائل"])
r += 1
for key, label, fmt in [("PM2", "سعر بيع المتر", NUM), ("PU", "سعر بيع الوحدة", NUM),
                        ("Net", "صافي متحصلات البيع", NUM), ("Profit", "صافي الربح", NUM),
                        ("ROI", "العائد الكلي", PCT), ("ROIa", "العائد السنوي", PCT),
                        ("ROIca", "العائد السنوي على النقد المدفوع فعلياً", PCT)]:
    ex.cell(r, 2, label)
    hl = key in ("Profit", "ROIa")
    style(ex.cell(r, 2), bold=hl, fill=GREEN if hl else None)
    for i, c in enumerate("CDE"):
        style(ex.cell(r, 3 + i, f"='الحسابات'!{c}{SC[key]}"), fmt=fmt, bold=hl, fill=GREEN if hl else None)
    r += 1
ex.cell(r, 2, "=\"سعر التعادل: \"&TEXT('الحسابات'!C" + str(SC["Break"]) + ",\"#,##0\")&\" ريال/م2 - أي أن السوق يجب أن ينخفض \"&TEXT(1-'الحسابات'!C" + str(SC["Break"]) + "/MktAvgM2,\"0%\")&\" عن متوسطه الحالي حتى يخسر المستثمر.\"")
ex.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
style(ex.cell(r, 2), color="7F7F7F", size=9, border=False)
r += 2

section(ex, r, "أثر شروط الصفقة على العائد السنوي (السيناريو الأساسي)", 4)
r += 1
header(ex, r, ["نسبة الدفع عند التعاقد", "المستثمر يتحمل الضريبة", "المطور يتحمل الضريبة"])
r += 1
net_base = f"'الحسابات'!D{SC['Net']}"
for down in (1, 0.5, 0.3):
    style(ex.cell(r, 2, down), fmt=PCT, bold=True)
    for i, rett in enumerate(("BuyCostRate", "0")):
        f = (f"=(1+({net_base}-DealValue*(1+{rett}))/(DealValue*{down}+DealValue*{rett}))^(12/Months)-1")
        style(ex.cell(r, 3 + i, f), fmt=PCT, fill=GREEN if (down == 0.3 and rett == "0") else None)
    r += 1
ex.cell(r, 2, "الدفع على مراحل (الباقي عند التسليم) وتحمّل المطور للضريبة يرفعان عائد المستثمر على النقد المدفوع بشكل كبير.")
ex.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
style(ex.cell(r, 2), color="7F7F7F", size=9, border=False)
r += 2

section(ex, r, "الخيار (2): الاحتفاظ والتأجير", 3)
r += 1
for label, f, fmt, hl in [
    ("الإيجار السنوي المتوقع للوحدة", "=Rent", NUM, False),
    ("صافي الإيجار السنوي (12 وحدة)", "=RentNet", NUM, False),
    ("العائد الإيجاري الصافي على سعرنا", "=YieldNet", PCT, True),
    ("العائد الإيجاري لو اشترى بسعر السوق", "=MktYield", PCT, False),
    ("عائد بديل (ودائع / صكوك)", "=BenchRate", PCT, False),
    ("فترة الاسترداد من الإيجار (سنة)", "=Payback", "0.0", False),
]:
    ex.cell(r, 2, label)
    style(ex.cell(r, 2), bold=hl, fill=GREEN if hl else None)
    ex.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    style(ex.cell(r, 3, f), fmt=fmt, bold=True, fill=GREEN if hl else None)
    style(ex.cell(r, 4), fill=GREEN if hl else None)
    r += 1

r += 1
section(ex, r, "نقاط القوة للمستثمر", 5)
r += 1
POINTS = [
    "=\"سعر جملة \"&TEXT(OurPriceM2,\"#,##0\")&\" ريال/م2 مقابل متوسط \"&TEXT(MktAvgM2,\"#,##0\")&\" ريال/م2 في السوق المحيط (خصم \"&TEXT(Discount,\"0%\")&\").\"",
    "=\"سعر الوحدة \"&TEXT(UnitPrice,\"#,##0\")&\" ريال - قريب من سعر الشقة في نفس الحي (مشروع نورفا يبدأ من 1,125,000 ريال) مع مساحة واستقلالية التاون هاوس.\"",
    "=\"هامش أمان: حتى في السيناريو المتحفظ (بيع بأقل 15% من السوق) الربح \"&TEXT('الحسابات'!C" + str(SC["Profit"]) + ",\"#,##0\")&\" ريال.\"",
    "=\"مرونة في الخروج: إعادة البيع عند التسليم أو التأجير بعائد صافي \"&TEXT(YieldNet,\"0.0%\")&\" سنوياً.\"",
    "الطلب على التاون هاوس في شرق الرياض مرتفع كحل وسط بين الشقة والفيلا، والقادسية حي قائم بخدماته.",
    "البيع على الخارطة تحت إشراف الهيئة العامة للعقار (وافي) وحساب ضمان للمدفوعات.",
]
for p in POINTS:
    ex.cell(r, 2, p)
    ex.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    style(ex.cell(r, 2))
    ex.row_dimensions[r].height = 30
    r += 1
r += 1
ex.cell(r, 2, "الأسعار المقارنة أسعار طلب معلنة على المنصات العقارية وقد تختلف عن أسعار البيع الفعلية. التفاصيل والمصادر في ورقة \"المشاريع المحيطة\".")
ex.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
style(ex.cell(r, 2), color="7F7F7F", size=9, border=False)

# رسم مقارنة سعر المتر
cr = 5
ex.cell(cr - 1, 7, "سعر المتر (ريال)")
style(ex.cell(cr - 1, 7), bold=True, fill=LIGHT)
ex.cell(cr - 1, 8, "القيمة")
style(ex.cell(cr - 1, 8), bold=True, fill=LIGHT)
ex.column_dimensions["G"].width = 30
ex.column_dimensions["H"].width = 12
labels = [f"='المشاريع المحيطة'!B{first + i}&\" \"&'المشاريع المحيطة'!D{first + i}&\"م2\"" for i in range(len(COMPS))]
for i, lab in enumerate(labels):
    style(ex.cell(cr + i, 7, lab))
    style(ex.cell(cr + i, 8, f"='المشاريع المحيطة'!F{first + i}"), fmt=NUM)
style(ex.cell(cr + len(COMPS), 7, "مشروعنا"), bold=True, fill=GREEN)
style(ex.cell(cr + len(COMPS), 8, "=OurPriceM2"), bold=True, fill=GREEN, fmt=NUM)
ch = BarChart()
ch.type = "bar"
ch.title = "سعر المتر: مشروعنا مقابل السوق"
ch.legend = None
ch.add_data(Reference(ex, min_col=8, min_row=cr - 1, max_row=cr + len(COMPS)), titles_from_data=True)
ch.set_categories(Reference(ex, min_col=7, min_row=cr, max_row=cr + len(COMPS)))
ch.series[0].graphicalProperties.solidFill = TEAL
ch.height, ch.width = 9, 14
ex.add_chart(ch, "G15")

order = ["عرض المستثمر", "المشاريع المحيطة", "الحسابات", "المدخلات"]
wb._sheets = [wb[n] for n in order]
wb.active = 0
for ws in wb:
    ws.sheet_properties.tabColor = TEAL if ws.title in order[:2] else "808080"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print("saved", OUT)
